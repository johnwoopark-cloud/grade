from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import sys

OUT = sys.argv[1]
F = 'Arial'
font = lambda **k: Font(name=F, size=k.pop('size', 10), **k)
INPUT = PatternFill('solid', fgColor='FFF2CC')
HEAD = PatternFill('solid', fgColor='DCE6F2')
ABSENT = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = '0000FF'

wb = Workbook()

# ---------- 설정 ----------
cfg = wb.active
cfg.title = '설정'
cfg['A1'] = '설정값 (노란 칸만 바꾸세요)'; cfg['A1'].font = font(bold=True, size=12)
rows = [
    ('질병 결시 인정 비율', 0.8, '학업성적관리규정 제22조 ④ (사용자 확인: 질병 결시 80%)'),
    ('출석인정 결시 인정 비율', 1.0, '학업성적관리규정 제22조 ⑥ 예시 (출석인정 100%)'),
    ('수행평가 결시 기본점수 (만점 대비)', 0.3, '사용자 지정: 수행평가 결시는 만점의 30%만 부여'),
]
cfg['A2'], cfg['B2'], cfg['C2'] = '항목', '값', '근거'
for c in 'ABC':
    cfg[f'{c}2'].font = font(bold=True); cfg[f'{c}2'].fill = HEAD; cfg[f'{c}2'].border = box
for i, (k, v, note) in enumerate(rows, start=3):
    cfg[f'A{i}'] = k; cfg[f'B{i}'] = v; cfg[f'C{i}'] = note
    cfg[f'A{i}'].font = font(); cfg[f'C{i}'].font = font(color='595959')
    cfg[f'B{i}'].font = font(color=BLUE); cfg[f'B{i}'].fill = INPUT; cfg[f'B{i}'].number_format = '0%'
    for c in 'ABC': cfg[f'{c}{i}'].border = box
cfg.column_dimensions['A'].width = 34; cfg.column_dimensions['B'].width = 10; cfg.column_dimensions['C'].width = 52
SICK, APPROVED, BASEPCT = '설정!$B$3', '설정!$B$4', '설정!$B$5'

# ---------- 계산기 ----------
ws = wb.create_sheet('계산기', 0)
ws['A1'] = '결시 인정점 계산기'; ws['A1'].font = font(bold=True, size=14)
ws['A2'] = ('노란 칸만 입력하세요. 보지 못한 항목은 \'응시\'를 \'결시\'로 바꾸고 사유를 고릅니다. '
            '정기시험 결시는 같은 과목에서 본 시험의 점수·평균으로(⑥-1), 정기시험이 모두 없으면 수행평가로(⑥-2) 계산합니다. '
            '수행평가 결시는 만점의 30%입니다.')
ws['A2'].font = font(color='595959'); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:M2'); ws.row_dimensions[2].height = 42
ws['A3'] = '평균: 결시한 시험은 그 시험의 과목평균, 응시한 시험은 그 시험의 과목평균(결시 계산에 쓰임)을 넣습니다. 위원회 심의(점수가 하나도 없음)일 때는 \'직접 입력\'에 학교가 정한 값을 넣습니다.'
ws['A3'].font = font(color='595959'); ws.merge_cells('A3:M3')

H = ['과목', '항목', '구분', '반영비율(%)', '만점', '응시/결시', '결시 사유', '원점수', '과목·영역 평균',
     '직접 입력(위원회)', '인정점', '적용 점수', '반영점수', '산출 근거']
HR = 5
for j, h in enumerate(H, start=1):
    c = ws.cell(HR, j, h); c.font = font(bold=True); c.fill = HEAD; c.border = box
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[HR].height = 30

subjects = []
for s in ['국어', '영어', '수학', '사회', '과학']:
    subjects.append((s, [('중간고사', '정기시험', 30, 100), ('기말고사', '정기시험', 30, 100),
                         ('수행평가1', '수행평가', 20, 20), ('수행평가2', '수행평가', 20, 20)]))
subjects.append(('과학실험', [('수행평가1', '수행평가', 50, 50), ('수행평가2', '수행평가', 50, 50)]))
items = [(s, *it) for s, its in subjects for it in its]
SPARE = 16
first = HR + 1
last = first + len(items) + SPARE - 1
rng = lambda col: f'${col}${first}:${col}${last}'

for k in range(len(items) + SPARE):
    r = first + k
    if k < len(items):
        s, name, typ, w, mx = items[k]
        ws[f'A{r}'], ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'E{r}'] = s, name, typ, w, mx
        ws[f'F{r}'] = '응시'; ws[f'G{r}'] = '질병'
    ratio = f'IF(G{r}="출석인정",{APPROVED},{SICK})'
    crit_exam = f'{rng("A")},A{r},{rng("C")},"정기시험",{rng("F")},"응시",{rng("H")},"<>"'
    crit_perf = f'{rng("A")},A{r},{rng("C")},"수행평가",{rng("F")},"응시",{rng("H")},"<>"'
    n_exam = f'COUNTIFS({crit_exam})'
    n_perf = f'COUNTIFS({crit_perf})'
    by = lambda crit: f'MIN(E{r},ROUND(SUMIFS({rng("H")},{crit})*I{r}/SUMIFS({rng("I")},{crit})*{ratio},2))'
    ws[f'K{r}'] = (f'=IF(F{r}<>"결시","",IF(J{r}<>"",J{r},'
                   f'IF(C{r}="수행평가",ROUND(E{r}*{BASEPCT},2),'
                   f'IF(I{r}="","평균 입력 필요",'
                   f'IFERROR(IF({n_exam}>0,{by(crit_exam)},IF({n_perf}>0,{by(crit_perf)},"위원회 심의")),"기준 시험 평균 입력 필요")))))')
    ws[f'L{r}'] = f'=IF(A{r}="","",IF(F{r}="결시",IF(ISNUMBER(K{r}),K{r},""),IF(H{r}="","",H{r})))'
    ws[f'M{r}'] = f'=IF(AND(ISNUMBER(L{r}),N(E{r})>0),L{r}/E{r}*D{r},"")'
    ws[f'N{r}'] = (f'=IF(F{r}<>"결시","",IF(J{r}<>"","직접 입력",'
                   f'IF(C{r}="수행평가","만점의 "&TEXT({BASEPCT},"0%")&" (기본점수)",'
                   f'IF({n_exam}>0,"⑥-1 응시한 정기시험 기준 × "&TEXT({ratio},"0%"),'
                   f'IF({n_perf}>0,"⑥-2 응시한 수행평가 기준 × "&TEXT({ratio},"0%"),"⑧ 위원회 심의")))))')
    for j in range(1, 15):
        c = ws.cell(r, j); c.border = box; c.font = font()
    for col in 'ABCDEFGHIJ':
        ws[f'{col}{r}'].fill = INPUT; ws[f'{col}{r}'].font = font(color=BLUE)
    for col in 'KLM': ws[f'{col}{r}'].number_format = '0.00'
    ws[f'N{r}'].font = font(color='595959', size=9)

dv = lambda opts: DataValidation(type='list', formula1='"' + ','.join(opts) + '"', allow_blank=True)
for opts, col in ((['응시', '결시'], 'F'), (['질병', '출석인정'], 'G'), (['정기시험', '수행평가'], 'C')):
    v = dv(opts); ws.add_data_validation(v); v.add(f'{col}{first}:{col}{last}')
ws.conditional_formatting.add(f'A{first}:N{last}', FormulaRule(formula=[f'$F{first}="결시"'], fill=ABSENT))

widths = [9, 11, 10, 9, 7, 9, 9, 9, 11, 11, 9, 9, 9, 34]
for j, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = f'C{first}'

# ---------- 요약 ----------
sm = wb.create_sheet('요약', 0)
sm['A1'] = '과목별 최종 점수'; sm['A1'].font = font(bold=True, size=14)
sm['A2'] = '계산기 시트에서 자동으로 계산됩니다. 과목을 추가하면 아래 빈 칸에 과목 이름만 적으세요.'; sm['A2'].font = font(color='595959')
hdr = ['과목', '최종 점수', '반영비율 합(%)', '결시 항목 수', '상태']
for j, h in enumerate(hdr, start=1):
    c = sm.cell(4, j, h); c.font = font(bold=True); c.fill = HEAD; c.border = box; c.alignment = Alignment(horizontal='center')
CR = lambda col: f'계산기!${col}${first}:${col}${last}'
names = [s for s, _ in subjects] + ['', '', '']
for i, nm in enumerate(names):
    r = 5 + i
    sm[f'A{r}'] = nm or None
    sm[f'B{r}'] = f'=IF(A{r}="","",ROUND(SUMIFS({CR("M")},{CR("A")},A{r}),2))'
    sm[f'C{r}'] = f'=IF(A{r}="","",SUMIFS({CR("D")},{CR("A")},A{r}))'
    sm[f'D{r}'] = f'=IF(A{r}="","",COUNTIFS({CR("A")},A{r},{CR("F")},"결시"))'
    sm[f'E{r}'] = (f'=IF(A{r}="","",IF(C{r}<>100,"반영비율 합이 100이 아님",'
                   f'IF(COUNTIFS({CR("A")},A{r})=0,"계산기에 항목 없음",'
                   f'IF(COUNTIFS({CR("A")},A{r},{CR("M")},"")>0,"입력 중","완료"))))')
    for j in range(1, 6):
        c = sm.cell(r, j); c.border = box; c.font = font()
    sm[f'A{r}'].fill = INPUT; sm[f'A{r}'].font = font(color=BLUE)
    sm[f'B{r}'].number_format = '0.00'; sm[f'B{r}'].font = font(bold=True)
for col, w in zip('ABCDE', [12, 12, 15, 12, 24]): sm.column_dimensions[col].width = w

# ---------- 사용법 ----------
hp = wb.create_sheet('사용법')
lines = [
    ('사용법', True),
    ('1. 계산기 시트의 노란 칸만 입력합니다. 흰 칸(인정점·적용 점수·반영점수·산출 근거)은 자동 계산됩니다.', False),
    ('2. 본 시험: 원점수와 그 시험의 과목평균을 넣습니다(과목평균은 다른 시험이 결시일 때 계산에 쓰임).', False),
    ('3. 못 본 시험: 응시/결시를 \'결시\'로, 결시 사유를 고르고, 그 시험의 과목평균을 넣습니다.', False),
    ('4. 수행평가 결시는 만점의 30%가 자동으로 들어갑니다(설정 시트에서 비율 변경).', False),
    ('5. 과목을 추가하려면 계산기 시트 아래 빈 줄에 과목·항목·구분·반영비율·만점을 적고, 요약 시트에 과목 이름을 적습니다.', False),
    ('', False),
    ('예시 (규정 캡처의 예시: 국어 기말 질병 결시)', True),
]
for i, (t, b) in enumerate(lines, start=1):
    hp[f'A{i}'] = t; hp[f'A{i}'].font = font(bold=b, size=12 if b and i == 1 else 10)
ex_h = ['항목', '응시/결시', '결시 사유', '원점수', '과목평균', '인정점']
ex = [('중간고사', '응시', '', 66.9, 68.72, ''), ('기말고사', '결시', '질병', '', 61.45, '47.86 (= 66.9 × 61.45 ÷ 68.72 × 80%)')]
for j, h in enumerate(ex_h, start=1):
    c = hp.cell(9, j, h); c.font = font(bold=True); c.fill = HEAD; c.border = box
for i, row in enumerate(ex, start=10):
    for j, v in enumerate(row, start=1):
        c = hp.cell(i, j, v if v != '' else None); c.font = font(); c.border = box
hp.column_dimensions['A'].width = 12
for col in 'BCDE': hp.column_dimensions[col].width = 11
hp.column_dimensions['F'].width = 38

wb.save(OUT)
print('saved', OUT, 'rows', first, last)
